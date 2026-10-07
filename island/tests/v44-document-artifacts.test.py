import os,sys,json,uuid,hashlib
from pathlib import Path
sys.path.insert(0,str(Path('server').resolve()))
from document_worker import run,read_document
from document_artifacts import capture_saved
from steward_host import HostReceipts
from docx import Document
from openpyxl import Workbook,load_workbook
root=Path('qa/v44')/('documents-'+uuid.uuid4().hex);root.mkdir(parents=True);art=(root/'artifacts').resolve()
os.environ.update(TERMINAL_CWD=str(root.resolve()),HD_ARTIFACT_DIR=str(art),HD_DOCUMENT_MODE='manual',HD_DOCUMENT_THEME='pixel',HD_DOCUMENT_LEDGER_ID='fixture-manual',HD_DOCUMENT_RUN_ID='fixture-provider',HD_DOCUMENT_PROJECT_ID='fixture-project')
checks=[]
def operation(op,path,**args):return run({'operation':op,'args':{'path':str(path.resolve()),**args}})
def snapshot(result,expected):
 a=result['artifact'];m=json.loads((art/'captures'/(a['id']+'.json')).read_text(encoding='utf-8'));p=art/'captures'/(a['id']+'.'+m['format'])
 assert m['sha256']==hashlib.sha256(p.read_bytes()).hexdigest()==a['sha256'];assert m['ledgerRunId']=='fixture-manual' and m['projectId']=='fixture-project' and m['mode']=='manual'
 r=operation('snapshot_read',p,sha256=m['sha256']);assert expected in r['content'];return p,m
text=root/'成果.txt';r=operation('write',text,content='版本一\n'+'\n'.join('行'+str(i) for i in range(250)));one,m=snapshot(r,'版本一')
assert r['verified'];r=operation('edit',text,oldText='版本一',newText='版本二');two,n=snapshot(r,'版本二');assert '版本一' in Path(r['backup']).read_text(encoding='utf-8');assert '版本一' in one.read_text(encoding='utf-8');checks.append('immutable versions and backup')
first=operation('snapshot_read',one,sha256=m['sha256']);second=operation('snapshot_read',one,sha256=m['sha256'],offset=first['nextOffset']);assert first['nextOffset']==121 and second['offset']==121 and first['totalLines']==251;checks.append('actual paginated snapshot text')
word=root/'排版.docx';doc=Document();p=doc.add_paragraph();p.add_run('金额：').bold=True;p.add_run('10').italic=True;p.add_run('0');doc.add_table(rows=1,cols=1).cell(0,0).text='100';doc.save(word)
r=operation('edit',word,oldText='100',newText='180');p,m=snapshot(r,'金额：180');updated=Document(word);assert updated.paragraphs[0].runs[0].bold and updated.paragraphs[0].runs[1].italic;assert updated.tables[0].cell(0,0).text=='180';checks.append('Word original formatting, table and actual preview')
excel=root/'预算.xlsx';book=Workbook();book.active['A1']='预算';book.active['B1']=100;book.active['C1']='=B1*2';book.save(excel);book.close()
r=operation('edit',excel,sheet='Sheet',cells=[{'address':'B1','value':180}]);p,m=snapshot(r,'C1==B1*2');book=load_workbook(excel);assert book.active['B1'].value==180 and book.active['C1'].value=='=B1*2';book.close();checks.append('Excel formula and original binary')
before=list((art/'captures').glob('*.json'));original=text.read_bytes()
try:operation('edit',text,oldText='不存在',newText='错误')
except ValueError:pass
else:raise AssertionError('invalid edit succeeded')
assert before==list((art/'captures').glob('*.json')) and text.read_bytes()==original;checks.append('failed edits create no artifact')
try:operation('edit',one,oldText='版本一',newText='篡改')
except ValueError:pass
else:raise AssertionError('managed snapshot editable')
try:operation('snapshot_read',text,sha256=hashlib.sha256(text.read_bytes()).hexdigest())
except ValueError:pass
else:raise AssertionError('unmanaged preview accepted')
try:operation('snapshot_read',one,sha256='0'*64)
except ValueError:pass
else:raise AssertionError('invalid snapshot checksum accepted')
checks.append('snapshot path and hash guarded')
try:capture_saved(text,'document_edit',1,expected_hash='0'*64)
except RuntimeError:pass
else:raise AssertionError('changed source registered')
checks.append('source-to-capture race rejected')
os.environ['HD_DOCUMENT_MODE']='island'
try:capture_saved(text,'document_write',1)
except RuntimeError:pass
else:raise AssertionError('automatic capture allowed')
os.environ['HD_DOCUMENT_MODE']='manual';checks.append('manual capture only')
receipt=HostReceipts();receipt.complete('unknown','document_write',{'path':str(text)},'not JSON');receipt.complete('failed','document_edit',{},json.dumps({'success':False,'error':'missing'}));receipt.complete('verified','document_edit',{},json.dumps(r));items=receipt.export();assert [i['status'] for i in items]==['unconfirmed','failed','done'];assert items[2]['artifact']['id']==r['artifact']['id'];checks.append('tool receipt is never fabricated from unknown text')
report={'passed':True,'scope':'Isolated local TXT, DOCX, XLSX tool calls and immutable snapshots, no model provider or user files.','directory':str(root),'checks':checks}
Path('qa/v44/document-artifacts-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps({'passed':True,'checks':len(checks)},ensure_ascii=False))
