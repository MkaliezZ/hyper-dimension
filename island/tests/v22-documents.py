import json,sys,os,uuid
from pathlib import Path
sys.path.insert(0,str(Path('server').resolve()))
from document_worker import run,read_document
from docx import Document
from openpyxl import Workbook,load_workbook
root=Path('qa/v22/documents')/uuid.uuid4().hex;root.mkdir(parents=True)
os.environ['TERMINAL_CWD']=str(root.resolve())
results=[]
text=root/'办公测试.txt';text.write_text('活动日期：10月10日\n预算：100\n',encoding='utf-8')
r=run({'operation':'edit','args':{'path':str(text.resolve()),'oldText':'10月10日','newText':'10月11日'}})
assert r['verified'] and '10月11日' in text.read_text(encoding='utf-8')
assert Path(r['backup']).read_text(encoding='utf-8').find('10月10日')>=0;results.append(r)
word=root/'办公测试.docx';doc=Document();paragraph=doc.add_paragraph()
paragraph.add_run('活动预算：').bold=True;paragraph.add_run('10').italic=True;paragraph.add_run('0；100')
doc.save(word)
r=run({'operation':'edit','args':{'path':str(word.resolve()),'oldText':'100','newText':'1000'}})
assert read_document(word).strip()=='活动预算：1000；1000'
updated=Document(word);assert updated.paragraphs[0].runs[0].bold and updated.paragraphs[0].runs[1].italic
assert '100；100' in read_document(Path(r['backup']));results.append(r)
excel=root/'办公测试.xlsx';book=Workbook();book.active['A1']='预算';book.active['B1']=100;book.active['C1']='=B1*2';book.save(excel);book.close()
r=run({'operation':'edit','args':{'path':str(excel.resolve()),'sheet':'Sheet','cells':[{'address':'B1','value':180}]}})
book=load_workbook(excel);assert book.active['B1'].value==180 and book.active['C1'].value=='=B1*2';book.close();results.append(r)
try:run({'operation':'write','args':{'path':str((root/'arbitrary.py').resolve()),'content':'pass'}})
except ValueError:pass
else:raise AssertionError('Document-only interface must reject executable source files')
try:run({'operation':'edit','args':{'path':str(text.resolve()),'oldText':'absent','newText':'x'}})
except ValueError:pass
else:raise AssertionError('Missing source text must not alter a document')
assert run({'operation':'list','args':{'path':str(root.resolve()),'pattern':'*.docx'}})['items']
Path('qa/v22/document-tools.json').write_text(json.dumps({'checks':7,'results':results},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'passed':7,'formats':['txt','docx','xlsx'],'backups':True,'formattingPreserved':True},ensure_ascii=False))

