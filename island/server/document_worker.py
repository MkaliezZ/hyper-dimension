"""Fixed document operations. No shell commands or model-authored Python are executed."""
import json,sys,os,uuid,shutil,re,hashlib
from document_artifacts import capture_saved
from pathlib import Path
TEXT={'.txt','.md','.csv','.json'}
FORMATS=TEXT|{'.docx','.xlsx','.pdf'}
MAX_BYTES=25*1024*1024

def scope_path(path):
    root=os.environ.get('HD_DOCUMENT_ROOT')
    if root:
        boundary=Path(root).resolve()
        resolved=path.resolve()
        if resolved!=boundary and boundary not in resolved.parents:
            raise ValueError('此路径不在你的管家文档工作区内；其他设备文件需要绑定本人执行端')
    return path

def document_path(raw,write=False):
    p=Path(raw).expanduser()
    if not p.is_absolute():p=Path(os.environ.get('TERMINAL_CWD',os.getcwd()))/p
    p=scope_path(p.resolve())
    managed=os.environ.get('HD_ARTIFACT_DIR')
    if managed and (p==Path(managed).resolve() or Path(managed).resolve() in p.parents):raise ValueError('成果快照与索引由手账管理，请操作原文档')
    if p.suffix.lower() not in FORMATS:raise ValueError('仅支持 txt、md、csv、json、docx、xlsx、pdf 文档')
    if write and p.suffix.lower()=='.pdf':raise ValueError('PDF 支持读取；请将编辑结果保存为 Word 或文本')
    if any(x.lower() in {'.ssh','.aws','.codex','.agents','.git','hermes-island'} for x in p.parts):raise ValueError('此目录不是普通文档目录')
    if p.exists() and p.stat().st_size>MAX_BYTES:raise ValueError('文档超过 25 MB，请拆分后处理')
    return p

def read_text(path):
    data=path.read_bytes()
    for encoding in ['utf-8-sig','gb18030']:
        try:return data.decode(encoding)
        except UnicodeDecodeError:pass
    raise ValueError('无法识别文本编码')

def docx_paragraphs(doc):
    def table_rows(table):
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
                for nested in cell.tables:yield from table_rows(nested)
    def all_paragraphs():
        yield from doc.paragraphs
        for table in doc.tables:yield from table_rows(table)
        for section in doc.sections:
            yield from section.header.paragraphs
            yield from section.footer.paragraphs
    seen=set()
    for paragraph in all_paragraphs():
        if paragraph._p not in seen:
            seen.add(paragraph._p);yield paragraph

def paragraphs_replace(paragraph,old,new):
    # Work backwards through matches so run offsets remain valid and newly
    # inserted text is never matched again. Refuse unsupported hyperlink spans.
    text=paragraph.text
    if ''.join(run.text for run in paragraph.runs)!=text:return 0
    starts=[m.start() for m in re.finditer(re.escape(old),text)]
    for start in reversed(starts):
        end=start+len(old);offset=0;began=False
        for run in paragraph.runs:
            value=run.text;left=offset;right=offset+len(value);offset=right
            if right<=start or left>=end:continue
            prefix=value[:max(0,start-left)];suffix=value[max(0,end-left):] if end<right else ''
            run.text=prefix+(new if not began else '')+suffix;began=True
    return len(starts)

def atomic_save(path,write):
    temporary=path.with_name('.'+path.stem+'.'+uuid.uuid4().hex+path.suffix)
    try:
        write(temporary)
        expected_hash=hashlib.sha256(temporary.read_bytes()).hexdigest()
        os.replace(temporary,path)
        return expected_hash
    finally:
        if temporary.exists():temporary.unlink()

def read_document(path):
    ext=path.suffix.lower()
    if ext in TEXT:return read_text(path)
    if ext=='.docx':
        from docx import Document
        return '\n'.join(p.text for p in docx_paragraphs(Document(path)))
    if ext=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=False,keep_links=False);out=[]
        try:
            for sheet in book:
                out.append('['+sheet.title+']')
                for row in sheet.iter_rows():
                    values=[f'{cell.coordinate}={cell.value}' for cell in row if cell.value is not None]
                    if values:out.append(' | '.join(values))
                    if len(out)>3000:out.append('[后续行已省略]');return '\n'.join(out)
            return '\n'.join(out)
        finally:book.close()
    from pypdf import PdfReader
    return '\n'.join((p.extract_text() or '') for p in PdfReader(path).pages[:100])

def backup(path,workdir):
    if not path.exists():return None
    target=Path(workdir)/'backups'/(uuid.uuid4().hex+'-'+path.name)
    target.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(path,target);return str(target)

def run(job):
    op=job['operation'];args=job.get('args') or {};workdir=os.environ.get('TERMINAL_CWD',os.getcwd())
    if op=='snapshot_read':
        root=Path(os.environ['HD_ARTIFACT_DIR']).resolve()/'captures';snapshot=Path(args.get('path','')).resolve()
        if snapshot.parent!=root or not re.fullmatch(r'capture-[a-f0-9]{32}\.(txt|md|csv|json|docx|xlsx|pdf)',snapshot.name):raise ValueError('无效成果快照')
        if not re.fullmatch(r'[a-f0-9]{64}',str(args.get('sha256',''))) or hashlib.sha256(snapshot.read_bytes()).hexdigest()!=args['sha256']:raise ValueError('成果快照校验失败')
        lines=read_document(snapshot).splitlines();offset=max(1,int(args.get('offset',1)));limit=min(120,max(1,int(args.get('limit',120))));end=min(len(lines),offset-1+limit)
        chunk='\n'.join(lines[offset-1:end]);shown=chunk[:16000]
        return {'content':shown,'offset':offset,'nextOffset':end+1 if end<len(lines) else None,'totalLines':len(lines),'truncated':len(chunk)>16000 or end<len(lines),'charactersTruncated':len(chunk)>16000}
    if op=='list':
        root=Path(args.get('path') or workdir).expanduser()
        if not root.is_absolute():root=Path(workdir)/root
        root=scope_path(root.resolve())
        if not root.is_dir():raise ValueError('目录不存在')
        pattern=str(args.get('pattern') or '*')
        if '/' in pattern or '\\' in pattern:raise ValueError('pattern 只能是文件名匹配')
        paths=root.rglob(pattern) if args.get('recursive') else root.glob(pattern);items=[]
        for index,p in enumerate(paths):
            if index>=5000:break
            try:scope_path(p)
            except ValueError:continue
            if any(x.lower() in {'.git','.ssh','.aws','.codex','.agents','node_modules'} for x in p.parts):continue
            if p.is_dir() and not args.get('recursive') or p.is_file() and p.suffix.lower() in FORMATS:
                items.append({'path':str(p),'directory':p.is_dir()})
            if len(items)>=100:break
        return {'items':items,'limited':len(items)>=100}
    path=document_path(args.get('path',''),write=op!='read')
    if op=='read':
        text=read_document(path).splitlines();offset=max(1,int(args.get('offset',1)));limit=min(300,max(1,int(args.get('limit',120))))
        chunk='\n'.join(f'{i+offset}|{line}' for i,line in enumerate(text[offset-1:offset-1+limit]))
        return {'path':str(path),'content':chunk[:16000],'totalLines':len(text),'truncated':len(chunk)>16000 or offset-1+limit<len(text)}
    path.parent.mkdir(parents=True,exist_ok=True);ext=path.suffix.lower();saved_backup=None;changes=0
    if op=='write':
        if path.exists():raise ValueError('文件已存在，请用 document_edit 修改，或提供新文件名')
        content=str(args.get('content') or '')
        if ext in TEXT:
            written_hash=atomic_save(path,lambda out:out.write_text(content,encoding='utf-8',newline=''));changes=len(content)
        elif ext=='.docx':
            from docx import Document
            doc=Document()
            for line in content.splitlines():doc.add_paragraph(line)
            written_hash=atomic_save(path,doc.save);changes=len(doc.paragraphs)
        elif ext=='.xlsx':
            from openpyxl import Workbook
            book=Workbook();sheet=book.active;sheet.title=str(args.get('sheet') or 'Sheet1')[:31]
            for row in args.get('rows') or []:sheet.append(row)
            written_hash=atomic_save(path,book.save);book.close();changes=sheet.max_row
    elif op=='edit':
        if not path.is_file():raise ValueError('文档不存在')
        old=str(args.get('oldText') or '');new=str(args.get('newText') or '')
        if ext in TEXT:
            text=read_text(path)
            if not old or old not in text:raise ValueError('未找到原文，请先读取文档确认')
            changes=text.count(old);new_content=text.replace(old,new);saved_backup=backup(path,workdir);written_hash=atomic_save(path,lambda out:out.write_text(new_content,encoding='utf-8',newline=''))
        elif ext=='.docx':
            from docx import Document
            if not old:raise ValueError('修改 Word 需要 oldText')
            doc=Document(path);changes=sum(paragraphs_replace(p,old,new) for p in docx_paragraphs(doc))
            if not changes:raise ValueError('未找到可编辑的原文，请先读取文档确认')
            saved_backup=backup(path,workdir);written_hash=atomic_save(path,doc.save)
        elif ext=='.xlsx':
            from openpyxl import load_workbook
            book=load_workbook(path)
            try:
                name=str(args.get('sheet') or book.active.title)
                if name not in book.sheetnames:raise ValueError('工作表不存在')
                sheet=book[name]
                for cell in args.get('cells') or []:
                    address=str(cell.get('address',''))
                    if not re.fullmatch(r'[A-Za-z]{1,3}[1-9][0-9]{0,6}',address):raise ValueError('单元格地址错误')
                    sheet[address]=cell.get('value');changes+=1
                if not changes:raise ValueError('没有需要修改的单元格')
                saved_backup=backup(path,workdir);written_hash=atomic_save(path,book.save)
            finally:book.close()
    else:raise ValueError('未知文档操作')
    # Re-open the saved file before returning success.
    verification=read_document(path)
    if hashlib.sha256(path.read_bytes()).hexdigest()!=written_hash:raise RuntimeError('保存后文件已被其他应用更改，请重新读取')
    artifact=None;artifact_error=None
    try:artifact=capture_saved(path,'document_'+op,changes,saved_backup,written_hash)
    except Exception as exc:artifact_error=str(exc)[:160]
    return {'success':True,'path':str(path),'backup':saved_backup,'changes':changes,'verified':True,'preview':verification[:1000],'artifact':artifact,'artifactError':artifact_error}

if __name__=='__main__':
    try:result=run(json.loads(sys.stdin.read()))
    except Exception as exc:result={'success':False,'error':type(exc).__name__+': '+str(exc)[:400]}
    print(json.dumps(result,ensure_ascii=False))

